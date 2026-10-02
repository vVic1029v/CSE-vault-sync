"""Sync cache -> Obsidian notes: one note per item, MOC indexes.

Entry kinds from sync ("module", "file", "page", "link", "locked") map to
notes here: contentless leaves become bare links in the parent's MOC so no
information is silently dropped. Child order is always Brightspace order.

EmitCtx bundles the per-course plumbing (vault paths, naming state) so the
emit_* functions stay small and never touch globals.
"""

import json
import re
from dataclasses import dataclass, field

from .config import AUTO_TAG, BASE  # noqa: F401  (installs vaultlib path)
from autolib import AUTO_END, AUTO_START, tags_line
from .htmlmd import html_to_md
from .utils import label, sanitize, vrel

TEXT_CAP = 100000  # max extracted chars per note; beyond that, see the file


@dataclass
class EmitCtx:
    """Everything emit_* needs besides the node itself."""
    vault: object       # Path: vault root (only output)
    vcourse: object     # Path: this course's folder in the vault
    raw_cdir: object    # Path: this course's folder in the RAW cache
    code: str           # sanitized course code (manifest key)
    cid: object         # Brightspace org-unit id
    used: dict = field(default_factory=dict)  # note_stem uniqueness state
    virt_folder: bool = False  # also emit VirtFolder `Folders:` parent links


def course_vault_name(raw_root, code):
    """Folder name for a course: `<PREFIX> - <short name>` (sanitized)."""
    info_path = raw_root / code / "course-info.json"
    name = ""
    if info_path.exists():
        try:
            name = json.loads(info_path.read_text(encoding="utf-8")).get("name", "")
        except Exception:
            pass
    short = re.sub(r"\s*\(20\d\d.*?\)\s*", "", name).strip()
    prefix = code.split("+")[0]
    # strip "CSE11B - ..." style prefix duplication only on a word boundary
    # (never eat single letters: "Bachelor..." must not match prefix "B").
    if len(prefix) > 1 and short.upper().startswith(prefix.upper() + " "):
        short = short[len(prefix):].strip(" -–:")
    return sanitize(f"{prefix} - {short}" if short else code, 80)


def wl(target, disp=None):
    """Obsidian path wikilink for target (pipe-safe)."""
    target = target.replace("|", "-")
    disp = label(disp or target.split("/")[-1])
    return f"[[{target}|{disp}]]" if disp != target.split("/")[-1] else f"[[{target}]]"


def moc_line(entry):
    """Render one MOC entry: note | file -> wikilink, bare -> markdown link."""
    if entry[0] in ("note", "file"):
        return f"- {wl(entry[1], entry[2])}"
    return f"- [{label(entry[2])}]({entry[1]})"


def viewer_url(cid, tid):
    """Brightspace viewer URL for a topic (used for bare/locked entries)."""
    return f"{BASE}/d2l/le/content/{cid}/viewContent/{tid}/View"


def extract_file_text(path):
    """Best-effort readable text of a file topic. '' when not extractable."""
    suf = path.suffix.lower()
    try:
        if suf == ".pdf":
            import pymupdf
            parts = []
            with pymupdf.open(path) as doc:
                for page in doc:
                    parts.append(page.get_text())
                    if sum(map(len, parts)) > TEXT_CAP:
                        break
            return re.sub(r"\n{3,}", "\n\n", "\n\n".join(parts)).strip()[:TEXT_CAP]
        if suf == ".pptx":
            from pptx import Presentation
            out = []
            for i, slide in enumerate(Presentation(str(path)).slides, 1):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text = shape.text.strip()
                        if text:
                            texts.append(text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            line = " | ".join(
                                c.text.strip() for c in row.cells).strip()
                            if line.strip("| "):
                                texts.append(line)
                if texts:
                    out.append(f"### Slide {i}\n\n" + "\n\n".join(texts))
            return "\n\n".join(out).strip()[:TEXT_CAP]
        if suf in (".txt", ".md", ".csv", ".s", ".py", ".c", ".h", ".java", ".vim"):
            return path.read_text(encoding="utf-8", errors="replace").strip()[:TEXT_CAP]
    except Exception:
        return ""
    return ""


def note_stem(vdir, title, used):
    """Deterministic unique note stem per directory (stable across rebuilds)."""
    base = sanitize(title or "untitled", 90) or "untitled"
    seen = used.setdefault(vdir.as_posix(), set())
    cand, i = base, 2
    while (cand + ".md").lower() in seen:
        cand = f"{base} ({i})"
        i += 1
    seen.add((cand + ".md").lower())
    return cand


def write_moc_note(path, title, up_target, brightspace_id, code,
                   contents, moc_entries, virt_folder=False):
    """Write one generated note: frontmatter + # Contents + # MOC.

    With virt_folder=True, also emit a VirtFolder `Folders:` parent link
    mirroring `up:` so the note nests in the VirtFolder tree.     Course-home notes are tree roots (their parent is the output folder's
    `! <folder>` top note, derived at build time).

    Frontmatter always starts at byte 0 (banner goes after the --- block):
    Obsidian/VirtFolder/Dataview only recognize leading frontmatter.
    Generated body is wrapped in AUTO markers (whole file is generated).
    """
    lines = ["---", f"up: [[{up_target}]]"]
    if virt_folder and up_target:
        lines += ["Folders:", f"  - \"[[{up_target}]]\""]
    lines += [tags_line(["brightspace"]),
              f"course_code: {code}", f"brightspace_id: {brightspace_id}",
              "---", AUTO_TAG.strip(), "", f"# {label(title)}", "",
              AUTO_START]
    if (contents or "").strip():
        lines += ["## Contents", "", contents.strip(), ""]
    if moc_entries:
        lines += ["## MOC", ""]
        lines += [moc_line(e) for e in moc_entries]
    lines += [AUTO_END]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def emit_topic(ctx, vdir, node, up_target):
    """One topic -> note, or a bare link entry when there is nothing to say."""
    title = node.get("title", "untitled") or "untitled"
    kind = node.get("kind")
    desc = (node.get("desc") or "").strip()
    if kind == "file":
        src = ctx.raw_cdir / node["path"]
        rel = vrel(ctx.vcourse / node["path"], ctx.vault)
        # PDFs are never converted to text: link them directly.
        text = "" if src.suffix.lower() == ".pdf" \
            else extract_file_text(src) if src.exists() else ""
        body = "\n\n".join(p for p in (desc, text) if p)
        if body:
            stem = note_stem(vdir, title, ctx.used)
            write_moc_note(vdir / f"{stem}.md", title, up_target,
                           node.get("id", ""), ctx.code, body,
                           [("file", rel, title)], ctx.virt_folder)
            return ("note", f"{vrel(vdir, ctx.vault)}/{stem}", title)
        return ("file", rel, title)
    if kind == "page":
        src = ctx.raw_cdir / node["path"]
        body = ""
        if src.exists():
            try:
                body = html_to_md(src.read_text(encoding="utf-8",
                                                errors="replace")).strip()
            except Exception:
                body = ""
        full = "\n\n".join(p for p in (desc, body) if p)
        if full:
            stem = note_stem(vdir, title, ctx.used)
            write_moc_note(vdir / f"{stem}.md", title, up_target,
                           node.get("id", ""), ctx.code, full, [], ctx.virt_folder)
            return ("note", f"{vrel(vdir, ctx.vault)}/{stem}", title)
        return ("bare", viewer_url(ctx.cid, node.get("id", "")), title)
    if kind == "link":
        url = node.get("url", "")
        full = "\n\n".join(p for p in (desc, f"[{label(title)}]({url})") if p)
        stem = note_stem(vdir, title, ctx.used)
        write_moc_note(vdir / f"{stem}.md", title, up_target,
                       node.get("id", ""), ctx.code, full, [], ctx.virt_folder)
        return ("note", f"{vrel(vdir, ctx.vault)}/{stem}", title)
    # locked / anything else: bare link, no note
    return ("bare", viewer_url(ctx.cid, node.get("id", "")), title)


def emit_module(ctx, vdir, node, up_target):
    """One module -> flat note in the parent dir + folder for its children.

    No content and no children -> bare link, nothing written.
    """
    title = node.get("title", "untitled") or "untitled"
    children = node.get("children", []) or []
    desc = (node.get("desc") or "").strip()
    if not desc and not children:
        return ("bare", f"{BASE}/d2l/le/content/{ctx.cid}/Home", title)
    stem = note_stem(vdir, title, ctx.used)
    subdir = vdir / node["name"]
    subdir.mkdir(parents=True, exist_ok=True)
    own = f"{vrel(vdir, ctx.vault)}/{stem}"
    kids = emit_level(ctx, subdir, own, children)
    write_moc_note(vdir / f"{stem}.md", title, up_target,
                   node.get("id", ""), ctx.code, desc, kids, ctx.virt_folder)
    return ("note", own, title)


def emit_level(ctx, vdir, up_target, nodes):
    """Emit one hierarchy level in Brightspace order. Never sorts."""
    entries = []
    for child in nodes or []:
        if child.get("kind") == "module":
            entries.append(emit_module(ctx, vdir, child, up_target))
        else:
            entries.append(emit_topic(ctx, vdir, child, up_target))
    return entries
